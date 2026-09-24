"""Add salary columns to the flips CSV and workbook.

Uses data/salary_history.json and data/seasons.csv. Safe to re-run: existing salary
columns are overwritten, not duplicated.
"""
import csv
import json
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CSV = ROOT / "mloty_stargard_teen_flips.csv"
XLSX = ROOT / "Mloty_Stargard_teen_flips.xlsx"

SAL_COLS = [
    "Salary in season bought ($/wk)",
    "Salary in season sold ($/wk)",
    "Buy ÷ salary",
    "Sale ÷ salary",
    "Salary at 22 ($/wk)",
    "Salary at 24 ($/wk)",
    "Salary at 26 ($/wk)",
    "Highest salary ($/wk)",
    "Last season with salary",
]

seasons = list(csv.DictReader(open(ROOT / "data" / "seasons.csv", encoding="utf-8")))
history = {pid: dict(map(tuple, h)) for pid, h in
           json.load(open(ROOT / "data" / "salary_history.json")).items() if h}


def season_of(day):
    for s in seasons:
        if s["start"][:10] <= day and (not s["finish"] or day < s["finish"][:10]):
            return int(s["season"])


def salary_fields(pid, age, bought, sold):
    h = {int(k): v for k, v in history[pid].items()}
    bs, ss = season_of(bought), season_of(sold)
    sb = h.get(bs)
    return [
        sb,
        h.get(ss),
        None,  # ratios: formulas in the workbook, computed in the CSV
        None,
        h.get(bs + 22 - age),
        h.get(bs + 24 - age),
        h.get(bs + 26 - age),
        max(h.values()),
        max(h),
    ]


# CSV
with open(CSV, encoding="utf-8-sig") as f:
    rows = list(csv.reader(f))
head = rows[0][:rows[0].index(SAL_COLS[0])] if SAL_COLS[0] in rows[0] else rows[0]
n = len(head)
idx = {c: i for i, c in enumerate(head)}
out = [head + SAL_COLS]
for r in rows[1:]:
    r = r[:n]
    v = salary_fields(r[idx["Player ID"]], int(r[idx["Age when bought"]]),
                      r[idx["Bought on"]], r[idx["Sold on"]])
    v[2] = round(float(r[idx["Buy price ($)"]]) / v[0], 1)
    v[3] = round(float(r[idx["Sale price ($)"]]) / v[0], 1)
    out.append(r + ["" if x is None else x for x in v])
with open(CSV, "w", encoding="utf-8", newline="") as f:
    csv.writer(f).writerows(out)

# Workbook
wb = openpyxl.load_workbook(XLSX)
ws = wb["Flips"]
first = n + 1  # first salary column
last_row = 1 + len(rows) - 1
hdr, money, ratio = ws.cell(1, n), ws.cell(2, 13), ws.cell(2, 17)  # R1, M2, Q2 as style templates
L = openpyxl.utils.get_column_letter
for j, name in enumerate(SAL_COLS):
    c = ws.cell(1, first + j, name)
    c.font, c.fill, c.alignment, c.border = copy(hdr.font), copy(hdr.fill), copy(hdr.alignment), copy(hdr.border)
    ws.column_dimensions[L(first + j)].width = 14
for i in range(2, last_row + 1):
    r = [ws.cell(i, k).value for k in range(1, n + 1)]
    v = salary_fields(str(r[idx["Player ID"]]), int(r[idx["Age when bought"]]),
                      r[idx["Bought on"]].strftime("%Y-%m-%d"), r[idx["Sold on"]].strftime("%Y-%m-%d"))
    sal = f"{L(first)}{i}"
    v[2] = f"=M{i}/{sal}"
    v[3] = f"=O{i}/{sal}"
    for j, x in enumerate(v):
        c = ws.cell(i, first + j, x)
        tmpl = ratio if j in (2, 3) else money if j < 8 else None
        if tmpl is not None:
            c.number_format, c.font = tmpl.number_format, copy(tmpl.font)
        else:
            c.font = copy(ws.cell(i, 1).font)
ws.auto_filter.ref = f"A1:{L(first + len(SAL_COLS) - 1)}{last_row}"

# Summary block beside the existing tables
sm = wb["Summary"]
F = L(first)
rng = lambda col: f"Flips!${col}$2:${col}${last_row}"
block = [
    ("Salary vs price", None),
    (None, None),
    ("Median salary in season bought ($/wk)", f"=MEDIAN({rng(F)})"),
    ("Players whose salary changed during the hold", f"=SUMPRODUCT(--({rng(F)}<>{rng(L(first + 1))}))"),
    ("Median buy ÷ salary", f"=MEDIAN({rng(L(first + 2))})"),
    ("Median sale ÷ salary", f"=MEDIAN({rng(L(first + 3))})"),
    ("Median highest salary, bought 2020 or earlier ($/wk)",
     f'=MEDIAN(IF(YEAR({rng("L")})<=2020,{rng(L(first + 7))}))'),
    ("Players ever reaching $50k/wk", f'=COUNTIF({rng(L(first + 7))},">=50000")'),
    ("Players ever reaching $100k/wk", f'=COUNTIF({rng(L(first + 7))},">=100000")'),
]
title_font, label_font = copy(sm["A1"].font), copy(sm["A3"].font)
for k, (label, formula) in enumerate(block):
    a, b = sm.cell(1 + k, 6, label), sm.cell(1 + k, 7)
    a.font = title_font if k == 0 else label_font
    if formula and formula.startswith("=MEDIAN(IF("):
        sm[b.coordinate] = openpyxl.worksheet.formula.ArrayFormula(b.coordinate, formula)
    elif formula:
        b.value = formula
    b.number_format = "0.0" if "÷" in (label or "") else "#,##0"
sm.column_dimensions["F"].width = 48
sm.column_dimensions["G"].width = 12

notes = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
note = ("Salary columns: weekly salary from each player's Salary History, matched to seasons by "
        "the official season dates; 'at 22/24/26' uses age when bought + seasons elapsed.")
if note not in notes:
    ws.cell(ws.max_row + 1, 1, note)
wb.save(XLSX)
print(f"added {len(SAL_COLS)} salary columns to {len(rows) - 1} players")
